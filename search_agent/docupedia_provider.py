# -*- coding: utf-8 -*-
"""
Konektor do wewnętrznej bazy wiedzy Bosch Docupedia (Atlassian Confluence).

Pozwala agentom na:
1. Wyszukiwanie dokumentacji technicznej, procedur i kart części za pomocą CQL (Confluence Query Language).
2. Pobieranie pełnej treści dokumentów wraz z tabelami i konwersję HTML na Markdown (przez markdownify).
"""
from __future__ import annotations

import logging
import os
import re
import sys
from pathlib import Path
from typing import Optional, Dict, Any, List, Union

logger = logging.getLogger(__name__)

from atlassian import Confluence
from markdownify import markdownify as md

from models import ProductQuery
from search_provider import BaseDatasheetProvider
from search_utils import generate_pn_variants, score_title_match, extract_relevant_lines

# Dołączenie katalogu agentic/ do ścieżki Pythona, aby wczytać config.yaml
_AGENTIC_DIR = Path(__file__).parent.parent / "agentic"
if str(_AGENTIC_DIR) not in sys.path:
    sys.path.append(str(_AGENTIC_DIR))

try:
    from llm_client import wczytaj_config
except ImportError:
    wczytaj_config = None

DEFAULT_DOCUPEDIA_URL = "https://inside-docupedia.bosch.com/confluence"


def get_docupedia_config() -> tuple[str, str]:
    """Zwraca (url, token) pobrane ze zmiennych środowiskowych lub config.yaml."""
    url = os.environ.get("DOCUPEDIA_URL", DEFAULT_DOCUPEDIA_URL)
    token = os.environ.get("DOCUPEDIA_PAT") or os.environ.get("DOCUPEDIA_TOKEN")

    if not token and wczytaj_config is not None:
        cfg_path = _AGENTIC_DIR / "config.yaml"
        if cfg_path.exists():
            try:
                cfg = wczytaj_config(cfg_path)
                token = cfg.get("docupedia_token") or cfg.get("ducopedia_token")
                url = cfg.get("docupedia_url", url)
            except Exception:
                pass

    return url, token or ""


class DocupediaClient:
    """Klient REST API dla Docupedii z obsługą konwersji Markdown."""

    def __init__(self, url: Optional[str] = None, token: Optional[str] = None):
        cfg_url, cfg_token = get_docupedia_config()
        self.url = (url or cfg_url).rstrip("/")
        self.token = token or cfg_token
        self._confluence = None

        if self.token:
            self._confluence = Confluence(
                url=self.url,
                token=self.token,
            )

    @property
    def is_configured(self) -> bool:
        return bool(self._confluence and self.token)

    def search_internal_knowledge_base(
        self,
        query: str,
        space_key: Optional[str] = None,
        limit: int = 5
    ) -> List[Dict[str, Any]]:
        """
        Wyszukuje artykuły i specyfikacje w Docupedii za pomocą CQL.
        Zwraca listę słowników: [{'id': str, 'title': str, 'space': str, 'url': str}].
        """
        if not self.is_configured:
            return []

        # Escapujemy zarówno backslash, jak i cudzysłów (backslash musi być
        # escapowany PIERWSZY, inaczej "\" na końcu zapytania mógłby
        # "połknąć" escapowanie następującego po nim cudzysłowu w CQL).
        clean_q = query.replace('\\', '\\\\').replace('"', '\\"')
        cql_parts = [f'text ~ "{clean_q}"']
        if space_key:
            cql_parts.append(f'space = "{space_key}"')

        cql = " AND ".join(cql_parts)
        try:
            results = self._confluence.cql(cql, limit=limit)
            items = []
            for item in results.get("results", []):
                content = item.get("content", {})
                page_id = content.get("id")
                title = content.get("title")
                space = content.get("space", {}).get("name", "Unknown")
                rel_url = content.get("_links", {}).get("webui", "")
                full_url = f"{self.url}{rel_url}" if rel_url else f"{self.url}/pages/viewpage.action?pageId={page_id}"

                items.append({
                    "id": page_id,
                    "title": title,
                    "type": content.get("type", "page"),
                    "space": space,
                    "url": full_url,
                })
            return items
        except Exception as e:
            print(f"[Docupedia] Błąd wyszukiwania CQL: {e}")
            return []

    def download_attachment(self, page_id: str, dest_dir: Path) -> Optional[str]:
        """Pobiera załącznik (plik) na dysk i zwraca bezwzględną ścieżkę."""
        if not self.is_configured:
            return None
        import requests
        try:
            meta = self._confluence.get_page_by_id(page_id, expand="version")
            if "_links" in meta and "download" in meta["_links"]:
                dl_url = self.url + meta["_links"]["download"]
                headers = {"Authorization": f"Bearer {self.token}"}
                resp = requests.get(dl_url, headers=headers, timeout=30)
                if resp.status_code == 200:
                    raw_name = meta.get("title", f"att_{page_id}.bin")
                    safe_name = re.sub(r'[\\/*?:"<>|]', "", raw_name).strip() or f"att_{page_id}.bin"
                    # WAŻNE: nazwę pliku poprzedzamy unikalnym ID strony/załącznika.
                    # Bez tego dwa różne wątki przetwarzające RÓŻNE części
                    # równolegle (tryb wsadowy CSV) mogły zapisać do TEGO SAMEGO
                    # pliku, jeśli oba załączniki nazywały się tak samo (np.
                    # "Datasheet.pdf") - jeden wątek potrafił wtedy odczytać do
                    # ekstrakcji zawartość podmienioną przez inny wątek w
                    # międzyczasie, czyli wymiary jednej części wyliczone z PDF-a
                    # zupełnie innej części.
                    filename = f"{page_id}_{safe_name}"
                    out_path = dest_dir / filename
                    out_path.write_bytes(resp.content)
                    return str(out_path.absolute())
        except Exception as e:
            print(f"[Docupedia] Błąd pobierania załącznika {page_id}: {e}")
        return None

    def get_internal_document_content(self, page_id: str, max_length: int = 15000) -> str:
        """
        Pobiera treść strony z Docupedii i konwertuje HTML na czytelny Markdown.
        Jeśli dokument jest długi, zamiast obcinać go od góry (co mogło ucinać
        tabelę wymiarów, jeśli znajdowała się dalej w treści), zachowujemy
        fragmenty w pobliżu wzmianek o wymiarach/wadze (patrz search_utils).
        """
        if not self.is_configured:
            return ""

        try:
            page = self._confluence.get_page_by_id(page_id, expand="body.storage")
            title = page.get("title", "Untitled")
            raw_html = page.get("body", {}).get("storage", {}).get("value", "")

            # Konwersja HTML do czystego Markdownu
            markdown_text = md(raw_html, strip=["script", "style"])
            # Usunięcie nadmiarowych pustych linii
            markdown_text = re.sub(r"\n{3,}", "\n\n", markdown_text).strip()

            if len(markdown_text) > max_length:
                lines = markdown_text.splitlines()
                approx_max_lines = max(150, max_length // 60)
                relevant_lines = extract_relevant_lines(lines, max_lines=approx_max_lines)
                markdown_text = "\n".join(relevant_lines)
                if len(markdown_text) > max_length:
                    markdown_text = markdown_text[:max_length]
                markdown_text += "\n\n... [Treść skrócona - zachowano fragmenty najbardziej związane z wymiarami/wagą] ..."

            return f"### Document: {title} (Docupedia ID: {page_id})\n\n{markdown_text}"
        except Exception as e:
            print(f"[Docupedia] Błąd pobierania strony {page_id}: {e}")
            return ""


def extract_text_from_attachment(file_path: Union[str, Path], part_number: str = "") -> str:
    """Wyciąga tekst z pobranych załączników PDF lub arkuszy Excel."""
    path = Path(file_path)
    if not path.exists():
        return ""

    ext = path.suffix.lower()
    pn_clean = re.sub(r"[\s\-_.]", "", part_number).lower() if part_number else ""
    keywords = [
        "mm", "cm", "kg", " g", "lbs", "oz", "weight", "mass", "length", "width", "height",
        "dimension", "size", "wymiar", "waga", "masa", "diameter", "thickness", "masse",
        "abmessung", "durchmesser", "gewicht", "gehäuse", "zeichnung", "drw", "dms"
    ]

    if ext == ".pdf":
        try:
            import pymupdf
            doc = pymupdf.open(str(path))
            extracted_lines = []
            seen = set()
            for page_num, page in enumerate(doc):
                text = page.get_text()
                for line in text.splitlines():
                    clean_l = line.strip()
                    if len(clean_l) < 3:
                        continue
                    lower_l = clean_l.lower()
                    has_kw = any(k in lower_l for k in keywords)
                    has_pn = bool(pn_clean and (pn_clean in re.sub(r"[\s\-_.]", "", lower_l)))
                    if has_kw or has_pn:
                        if clean_l not in seen:
                            seen.add(clean_l)
                            extracted_lines.append(f"[P.{page_num+1}] {clean_l}")
            if extracted_lines:
                return f"### PDF Attachment Content: {path.name}\n" + "\n".join(extracted_lines[:150])
        except Exception as e:
            logger.warning("Error extracting text from PDF %s: %s", path.name, e)

    elif ext in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
            extracted_rows = []
            for sheet_name in wb.sheetnames[:5]:
                ws = wb[sheet_name]
                row_count = 0
                for row in ws.iter_rows(values_only=True):
                    row_count += 1
                    if row_count > 1000:
                        break
                    row_vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if not row_vals:
                        continue
                    row_str = " | ".join(row_vals)
                    lower_r = row_str.lower()
                    has_kw = any(k in lower_r for k in keywords)
                    has_pn = bool(pn_clean and (pn_clean in re.sub(r"[\s\-_.]", "", lower_r)))
                    if has_pn or (has_kw and len(row_vals) <= 12):
                        extracted_rows.append(f"[{sheet_name}:R{row_count}] {row_str[:200]}")
                    if len(extracted_rows) >= 50:
                        break
                if len(extracted_rows) >= 50:
                    break
            if extracted_rows:
                return f"### Excel Attachment Content: {path.name}\n" + "\n".join(extracted_rows)
        except Exception as e:
            logger.warning("Error extracting text from Excel %s: %s", path.name, e)

    return ""


class DocupediaDatasheetProvider(BaseDatasheetProvider):
    """
    Dostawca kart katalogowych pobierający dane z wewnętrznej Docupedii Bosch.
    """

    def __init__(self, client: Optional[DocupediaClient] = None, max_pages_to_fetch: int = 3):
        self.client = client or DocupediaClient()
        self.max_pages = max_pages_to_fetch

    def search_and_fetch_text(self, query: ProductQuery) -> Dict[str, Any]:
        if not self.client.is_configured:
            return {
                "title": "Docupedia not configured",
                "source_url": "",
                "text": "Docupedia token is not configured in config.yaml.",
                "has_data": False,
                "source_type": "docupedia",
            }

        pn = str(query.part_number).strip()
        queries_to_try = generate_pn_variants(pn)
        if query.title:
            queries_to_try.append(f'"{pn}" {query.title}')
            queries_to_try.append(f"{pn} {query.title}")

        found_pages: List[Dict[str, Any]] = []
        seen_ids = set()
        pool_cap = max(self.max_pages * 3, 8)

        for q_term in queries_to_try:
            pages = self.client.search_internal_knowledge_base(q_term, limit=max(self.max_pages, 5))
            for p in pages:
                pid = p.get("id")
                if pid and pid not in seen_ids:
                    seen_ids.add(pid)
                    found_pages.append(p)
            if len(found_pages) >= pool_cap:
                break

        # Fallback poszukiwania specyfikacji rodziny części:
        # Jeśli szukana część to np. sensor (TORQUE SENSOR, DMS) lub ma prefiks 7802,
        # poszukaj specyfikacji technicznych / Pflichtenheft danej rodziny
        title_l = (query.title or "").lower()
        if len(found_pages) < pool_cap and ("torque" in title_l or "sensor" in title_l or pn.startswith("7802")):
            spec_cqls = [
                'text ~ "7802" AND (text ~ "Pflichtenheft" OR text ~ "SCS" OR text ~ "Consolidation")',
                'title ~ "Drehmomentsensor*" OR title ~ "Torque Sensor*"'
            ]
            for cql in spec_cqls:
                try:
                    cql_res = self.client._confluence.cql(cql, limit=3)
                    for item in cql_res.get("results", []):
                        c = item.get("content", {})
                        cid = c.get("id")
                        if cid and cid not in seen_ids:
                            seen_ids.add(cid)
                            rel_url = c.get("_links", {}).get("webui", "")
                            full_url = f"{self.client.url}{rel_url}" if rel_url else f"{self.client.url}/pages/viewpage.action?pageId={cid}"
                            found_pages.append({
                                "id": cid,
                                "title": c.get("title", ""),
                                "type": c.get("type", "page"),
                                "space": c.get("space", {}).get("name", "Unknown"),
                                "url": full_url,
                            })
                except Exception:
                    pass

        if not found_pages:
            return {
                "title": f"Docupedia search for {query.brand} {pn}",
                "source_url": f"{self.client.url}/dosearchsite.action?queryString={pn}",
                "text": f"No internal Docupedia documentation found for part number {pn}.",
                "has_data": False,
                "source_type": "docupedia",
            }

        # Strony, których TYTUŁ zawiera numer części lub słowa specyfikacji,
        # są oceniane wyżej (patrz score_title_match).
        found_pages.sort(key=lambda p: score_title_match(p.get("title", ""), queries_to_try), reverse=True)
        selected_pages = found_pages[: self.max_pages]

        # Pobranie treści dopasowanych stron i załączników
        docs_text = []
        attachments = []
        primary_url = selected_pages[0].get("url", "")
        primary_title = selected_pages[0].get("title", "")

        download_dir = Path(__file__).parent / "downloads"
        download_dir.mkdir(exist_ok=True)

        for p in selected_pages:
            p_id = str(p["id"])
            p_type = p.get("type", "page")
            p_title = str(p.get("title", ""))

            # Jeśli to załącznik (lub plik po tytule), pobierz go i wyciągnij tekst
            if p_type == "attachment" or p_title.lower().endswith((".pdf", ".xlsx", ".xls", ".png", ".jpg")):
                dl_path = self.client.download_attachment(p_id, download_dir)
                if dl_path:
                    attachments.append(dl_path)
                    att_text = extract_text_from_attachment(dl_path, part_number=pn)
                    if att_text:
                        docs_text.append(att_text)

                # Dla załącznika sprawdź również stronę nadrzędną (container page)
                try:
                    meta = self.client._confluence.get_page_by_id(p_id, expand="container")
                    container = meta.get("container", {})
                    c_id = container.get("id")
                    c_title = container.get("title")
                    if c_id:
                        c_text = self.client.get_internal_document_content(str(c_id), max_length=4000)
                        if c_text.strip():
                            docs_text.append(f"### Parent Container Page: {c_title} (ID: {c_id})\n\n{c_text}")
                except Exception:
                    pass

            # Pobierz także ewentualny tekst z Markdowna strony
            content = self.client.get_internal_document_content(p_id)
            if content.strip():
                docs_text.append(content)

        if not docs_text and not attachments:
            return {
                "title": primary_title,
                "source_url": primary_url,
                "text": f"Found pages in Docupedia, but their content could not be read for {pn}.",
                "has_data": False,
                "source_type": "docupedia",
                "attachments": [],
            }

        combined = "\n\n---\n\n".join(docs_text)
        return {
            "title": primary_title,
            "source_url": primary_url,
            "text": f"BOSCH DOCUPEDIA INTERNAL DOCUMENTATION FOR {query.brand} {pn} ({query.title or ''}):\n\n{combined}",
            "has_data": True,
            "source_type": "docupedia",
            "attachments": attachments,
        }
